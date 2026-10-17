"""Sarvam Check — grade handwritten answer sheets with Sarvam Vision + Sarvam LLM."""
import csv
import io
import json
import os
import threading
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import PurePosixPath

from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError

import sarvam_client

load_dotenv()

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 500 * 1024 * 1024

ALLOWED = (".pdf", ".png", ".jpg", ".jpeg")
JOBS = {}
JOBS_LOCK = threading.Lock()


def _split_pdf_pages(stem, name, data):
    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) <= 1:
        return [(stem, name, data)]
    out = []
    for i, page in enumerate(reader.pages, start=1):
        w = PdfWriter()
        w.add_page(page)
        buf = io.BytesIO()
        w.write(buf)
        out.append((f"{stem} (page {i})", f"{stem}_page_{i}.pdf", buf.getvalue()))
    return out


def expand_uploads(files, split_pages=False):
    """Turn uploaded PDFs/images/ZIPs into a flat list of (student_name, filename, bytes)."""
    raw = []
    for f in files:
        name = f.filename or "upload"
        data = f.read()
        if name.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                for info in zf.infolist():
                    inner = PurePosixPath(info.filename)
                    if info.is_dir() or inner.name.startswith((".", "__MACOSX")) or "__MACOSX" in inner.parts:
                        continue
                    if inner.suffix.lower() in ALLOWED:
                        raw.append((inner.stem, inner.name, zf.read(info)))
        elif name.lower().endswith(ALLOWED):
            raw.append((PurePosixPath(name).stem, name, data))

    if not split_pages:
        return raw
    subs = []
    for stem, name, data in raw:
        subs.extend(_split_pdf_pages(stem, name, data) if name.lower().endswith(".pdf") else [(stem, name, data)])
    return subs


def process_submission(job_id, idx, filename, data, scheme, detect_name):
    entry = JOBS[job_id]["results"][idx]
    try:
        entry["status"] = "reading"
        text = sarvam_client.ocr_document(filename, data, os.getenv("DOC_LANGUAGE", "en-IN"))
        entry["ocr_text"] = text
        entry["status"] = "grading"
        entry["grade"] = sarvam_client.grade_answer(scheme, text)
        if detect_name and entry["grade"]["student_name"]:
            entry["student"] = entry["grade"]["student_name"]
        entry["status"] = "done"
    except Exception as e:  # surface every failure in the UI rather than killing the batch
        entry["status"] = "error"
        entry["error"] = str(e)


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/config")
def config():
    key = os.getenv("SARVAM_API_KEY", "")
    return jsonify(
        key_configured=bool(key) and key != "your_sarvam_api_key_here",
        model=os.getenv("SARVAM_CHAT_MODEL", "sarvam-105b"),
        workers=int(os.getenv("MAX_WORKERS", "4")),
        language=os.getenv("DOC_LANGUAGE", "en-IN"),
    )


@app.post("/api/dry-check")
def dry_check():
    try:
        reply = sarvam_client.ping()
    except Exception as e:
        return jsonify(ok=False, error=str(e)), 502
    return jsonify(ok=True, reply=reply.strip())


@app.post("/api/jobs")
def create_job():
    try:
        scheme = json.loads(request.form.get("scheme", "{}"))
    except json.JSONDecodeError:
        return jsonify(error="Invalid marking scheme"), 400
    if not scheme.get("questions"):
        return jsonify(error="Marking scheme has no questions"), 400

    try:
        subs = expand_uploads(request.files.getlist("files"), request.form.get("split_pages") == "1")
    except zipfile.BadZipFile:
        return jsonify(error="One of the ZIP files is corrupted"), 400
    except PdfReadError:
        return jsonify(error="One of the PDFs could not be read"), 400
    detect_name = request.form.get("detect_name") == "1"
    if not subs:
        return jsonify(error="No PDF or image submissions found in the upload"), 400

    job_id = uuid.uuid4().hex[:12]
    JOBS[job_id] = {
        "scheme": scheme,
        "results": [
            {"student": student, "file": fname, "status": "queued", "grade": None, "error": None, "ocr_text": ""}
            for student, fname, _ in subs
        ],
    }
    pool = ThreadPoolExecutor(max_workers=int(os.getenv("MAX_WORKERS", "4")))
    for i, (_, fname, data) in enumerate(subs):
        pool.submit(process_submission, job_id, i, fname, data, scheme, detect_name)
    pool.shutdown(wait=False)
    return jsonify(job_id=job_id, count=len(subs))


@app.get("/api/jobs/<job_id>")
def job_status(job_id):
    job = JOBS.get(job_id)
    if not job:
        return jsonify(error="Job not found"), 404
    return jsonify(scheme=job["scheme"], results=job["results"])


def _rows(job):
    qids = [q["id"] for q in job["scheme"]["questions"]]
    max_by_q = {q["id"]: sum(float(p["marks"]) for p in q["points"]) for q in job["scheme"]["questions"]}
    max_total = sum(max_by_q.values())
    header = ["Student", "Roll no.", "File"] + [f"{q} (/{max_by_q[q]:g})" for q in qids] + [
        "Total", "Out of", "Percentage", "Overall feedback", "Per-question feedback", "Status"]
    rows = []
    for r in job["results"]:
        g = r.get("grade")
        if g:
            qs = {q["question_id"]: q for q in g["questions"]}
            per_q = [qs.get(q, {}).get("score", 0) for q in qids]
            fb = " | ".join(f"{q}: {qs.get(q, {}).get('feedback', '')}" for q in qids)
            pct = round(100 * g["total_score"] / max_total, 1) if max_total else 0
            rows.append([r["student"], g.get("roll_number", ""), r["file"], *per_q, g["total_score"], max_total,
                         pct, g["overall_feedback"], fb, "graded"])
        else:
            rows.append([r["student"], "", r["file"], *[""] * len(qids), "", max_total, "",
                         "", "", r.get("error") or r["status"]])
    return header, rows


@app.get("/api/jobs/<job_id>/export")
def export(job_id):
    job = JOBS.get(job_id)
    if not job:
        return jsonify(error="Job not found"), 404
    header, rows = _rows(job)

    if request.args.get("format") == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(rows)
        return send_file(io.BytesIO(buf.getvalue().encode("utf-8-sig")), mimetype="text/csv",
                         as_attachment=True, download_name="scores.csv")

    wb = Workbook()
    ws = wb.active
    ws.title = "Scores"
    ws.append(header)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    widths = {"Student": 22, "File": 26, "Overall feedback": 60, "Per-question feedback": 80}
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(h, 12)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="scores.xlsx")


if __name__ == "__main__":
    app.run(debug=True, port=int(os.getenv("PORT", "5000")), threaded=True)
